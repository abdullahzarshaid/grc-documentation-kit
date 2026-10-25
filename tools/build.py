"""Build the kit: Markdown masters in src/ -> Word documents in kit/<framework>/, a single
browsable page docs/kit.html, and the document index tables in README.md.

Usage: python tools/build.py            (build everything)
       python tools/build.py --index     (rewrite README index only)
Standard library plus python-docx, PyYAML and markdown-it-py.
"""
import html as html_mod
import json
import re
import sys
from datetime import date
from pathlib import Path

import yaml
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm
from markdown_it import MarkdownIt

ROOT = Path(__file__).resolve().parents[1]
SRC, KIT, DOCS = ROOT / "src", ROOT / "kit", ROOT / "docs"
FRAMEWORKS = {
    "iso27001": ("ISO/IEC 27001:2022", "Information security management system"),
    "iso27701": ("ISO/IEC 27701", "Privacy information management, the PIMS extension"),
    "soc2": ("SOC 2", "AICPA Trust Services Criteria: Security, Availability, Confidentiality, Processing Integrity, Privacy"),
    "gdpr": ("GDPR", "Operational privacy for a product or service that processes personal data"),
    "ai": ("AI governance", "Policies and registers for the use and development of AI systems"),
    "cross": ("Cross-framework", "Playbooks, crosswalks, audit tools and governance documents that serve every framework"),
}
TYPES = ["policy", "procedure", "plan", "statement", "record", "template", "checklist", "playbook", "guide"]
NAVY, SLATE, ACCENT = RGBColor(0x0B, 0x2D, 0x4D), RGBColor(0x52, 0x64, 0x78), RGBColor(0x16, 0x61, 0xA2)
md = MarkdownIt("commonmark").enable("table")


# ----------------------------------------------------------------------------- loading
def load_docs():
    docs = []
    for p in sorted(SRC.rglob("*.md")):
        if p.name == "README.md":
            continue
        text = p.read_text(encoding="utf-8")
        m = re.match(r"^---\n(.*?)\n---\n(.*)$", text, flags=re.S)
        if not m:
            raise SystemExit(f"{p}: missing front matter")
        meta = yaml.safe_load(m.group(1))
        meta["_path"], meta["_body"] = p, m.group(2).strip() + "\n"
        docs.append(meta)
    ids = [d["id"] for d in docs]
    dup = {i for i in ids if ids.count(i) > 1}
    if dup:
        raise SystemExit(f"duplicate ids: {dup}")
    return docs


# ----------------------------------------------------------------------------- docx helpers
def set_cell_shading(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def add_inline(par, text, bold=False, italic=False, size=None, color=None):
    """Render **bold**, *italic* and `code` inline markup inside a paragraph."""
    pos = 0
    for m in re.finditer(r"(\*\*(.+?)\*\*|\*(.+?)\*|`(.+?)`)", text):
        if m.start() > pos:
            r = par.add_run(text[pos:m.start()]); r.bold, r.italic = bold, italic
            if size: r.font.size = Pt(size)
            if color: r.font.color.rgb = color
        if m.group(2):
            r = par.add_run(m.group(2)); r.bold = True
        elif m.group(3):
            r = par.add_run(m.group(3)); r.italic = True
        else:
            r = par.add_run(m.group(4)); r.font.name = "Consolas"
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = color
        pos = m.end()
    if pos < len(text):
        r = par.add_run(text[pos:]); r.bold, r.italic = bold, italic
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = color


def style_document(doc):
    st = doc.styles["Normal"]; st.font.name, st.font.size = "Calibri", Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    st.paragraph_format.space_after = Pt(6)
    for name, size in (("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11.5)):
        h = doc.styles[name]; h.font.name, h.font.size, h.font.bold = "Calibri", Pt(size), True
        h.font.color.rgb = NAVY; h.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        h.paragraph_format.space_before, h.paragraph_format.space_after = Pt(14 if name == "Heading 1" else 10), Pt(4)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.2); s.top_margin = s.bottom_margin = Cm(2)


def add_footer(doc, meta):
    for s in doc.sections:
        p = s.footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"{meta['id']}  |  {meta['title']}  |  Version 1.0  |  {meta.get('classification', 'Internal')}  |  Page ")
        r.font.size, r.font.color.rgb = Pt(8), SLATE
        for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
            run = p.add_run(); run.font.size = Pt(8); run.font.color.rgb = SLATE
            if tag:
                el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tag); run._r.append(el)
            else:
                el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = txt; run._r.append(el)


def control_table(doc, meta):
    rows = [("Document ID", meta["id"]), ("Title", meta["title"]), ("Type", meta["type"].capitalize()),
            ("Framework", FRAMEWORKS[meta["framework"]][0]),
            ("Requirements served", "; ".join(meta.get("maps", []))),
            ("Owner (role)", meta["owner"]), ("Approver (role)", meta.get("approver", "")),
            ("Classification", meta.get("classification", "Internal")),
            ("Version", "1.0 (template)"), ("Effective date", "[Date]"), ("Review", meta.get("review", "Annual"))]
    t = doc.add_table(rows=0, cols=2); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for k, v in rows:
        c = t.add_row().cells
        c[0].text, c[1].text = k, v
        set_cell_shading(c[0], "EEF3F8")
        for cell in c:
            for par in cell.paragraphs:
                par.paragraph_format.space_after = Pt(2)
                for run in par.runs:
                    run.font.size = Pt(9.5)
        c[0].paragraphs[0].runs[0].bold = True
    t.columns[0].width = Cm(4.5); t.columns[1].width = Cm(12)
    for r in t.rows:
        r.cells[0].width, r.cells[1].width = Cm(4.5), Cm(12)
    doc.add_paragraph()
    h = doc.add_paragraph(); r = h.add_run("Revision history"); r.bold = True; r.font.color.rgb = NAVY
    t2 = doc.add_table(rows=2, cols=4); t2.style = "Table Grid"
    for i, k in enumerate(("Version", "Date", "Change", "Approved by")):
        t2.rows[0].cells[i].text = k; set_cell_shading(t2.rows[0].cells[i], "EEF3F8")
        t2.rows[0].cells[i].paragraphs[0].runs[0].bold = True
    for i, k in enumerate(("1.0", "[Date]", "First issue from the kit template", "[Role]")):
        t2.rows[1].cells[i].text = k
    for row in t2.rows:
        for cell in row.cells:
            for par in cell.paragraphs:
                par.paragraph_format.space_after = Pt(2)
                for run in par.runs:
                    run.font.size = Pt(9.5)
    doc.add_paragraph()


def render_tokens(doc, tokens):
    """Walk markdown-it tokens and write them into the docx."""
    i, n = 0, len(tokens)
    list_stack = []  # ("bullet"|"ordered", counter)
    while i < n:
        t = tokens[i]
        if t.type == "heading_open":
            level = int(t.tag[1])
            text = tokens[i + 1].content
            doc.add_heading(text, level=min(max(level - 1, 1), 3))
            i += 3; continue
        if t.type == "paragraph_open":
            content = tokens[i + 1].content
            p = doc.add_paragraph()
            if list_stack:
                kind = list_stack[-1][0]
                p.style = "List Bullet" if kind == "bullet" else "List Number"
                if len(list_stack) > 1:
                    p.paragraph_format.left_indent = Cm(1.2 * len(list_stack))
            add_inline(p, content.replace("\n", " "))
            i += 3; continue
        if t.type in ("bullet_list_open", "ordered_list_open"):
            list_stack.append(("bullet" if t.type == "bullet_list_open" else "ordered", 0)); i += 1; continue
        if t.type in ("bullet_list_close", "ordered_list_close"):
            list_stack.pop(); i += 1; continue
        if t.type == "table_open":
            # collect rows
            rows, j = [], i + 1
            cur = None
            while tokens[j].type != "table_close":
                if tokens[j].type == "tr_open":
                    cur = []
                elif tokens[j].type in ("th_open", "td_open"):
                    cur.append(tokens[j + 1].content)
                elif tokens[j].type == "tr_close":
                    rows.append(cur)
                j += 1
            tbl = doc.add_table(rows=0, cols=len(rows[0])); tbl.style = "Table Grid"
            for ri, row in enumerate(rows):
                cells = tbl.add_row().cells
                for ci, val in enumerate(row):
                    par = cells[ci].paragraphs[0]; par.paragraph_format.space_after = Pt(2)
                    add_inline(par, val, bold=(ri == 0), size=9.5)
                    if ri == 0:
                        set_cell_shading(cells[ci], "EEF3F8")
            doc.add_paragraph()
            i = j + 1; continue
        if t.type == "fence" or t.type == "code_block":
            for line in t.content.rstrip("\n").split("\n"):
                p = doc.add_paragraph(); r = p.add_run(line); r.font.name = "Consolas"; r.font.size = Pt(9)
                p.paragraph_format.space_after = Pt(0)
            doc.add_paragraph()
            i += 1; continue
        if t.type == "blockquote_open":
            j = i + 1
            while tokens[j].type != "blockquote_close":
                if tokens[j].type == "inline":
                    p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(1)
                    add_inline(p, tokens[j].content, italic=True, color=SLATE)
                j += 1
            i = j + 1; continue
        if t.type == "hr":
            doc.add_paragraph(); i += 1; continue
        i += 1


def build_docx(meta):
    doc = Document(); style_document(doc)
    p = doc.add_paragraph(); r = p.add_run(meta["id"]); r.font.size = Pt(10); r.font.color.rgb = ACCENT; r.bold = True
    p = doc.add_paragraph(); r = p.add_run(meta["title"]); r.font.size = Pt(22); r.bold = True; r.font.color.rgb = NAVY
    p = doc.add_paragraph(); r = p.add_run(meta.get("summary", "")); r.italic = True; r.font.color.rgb = SLATE
    control_table(doc, meta)
    render_tokens(doc, md.parse(meta["_body"]))
    doc.add_paragraph()
    p = doc.add_paragraph(); r = p.add_run("This document is a template from an open documentation kit. Replace every item in square brackets, remove what does not apply, and have the approver sign it before use. A template is not evidence of compliance.")
    r.font.size = Pt(8.5); r.font.color.rgb = SLATE; r.italic = True
    add_footer(doc, meta)
    out = KIT / meta["framework"] / (meta["_path"].stem + ".docx")
    out.parent.mkdir(parents=True, exist_ok=True); doc.save(out)
    return out


# ----------------------------------------------------------------------------- html site
def md_to_html(body):
    h = md.render(body)
    return re.sub(r"<table>", '<table class="t">', h)


def build_site(docs, registers):
    nav, body = [], []
    for fw, (name, desc) in FRAMEWORKS.items():
        group = [d for d in docs if d["framework"] == fw]
        regs = [r for r in registers if r["framework"] == fw]
        if not group and not regs:
            continue
        nav.append(f'<div class="g"><b>{html_mod.escape(name)}</b>')
        body.append(f'<section id="{fw}"><h1>{html_mod.escape(name)}</h1><p class="lead">{html_mod.escape(desc)}</p>')
        for d in group:
            anchor = d["id"].lower()
            nav.append(f'<a href="#{anchor}">{html_mod.escape(d["id"])} {html_mod.escape(d["title"])}</a>')
            maps = "; ".join(html_mod.escape(m) for m in d.get("maps", []))
            rel = f'kit/{fw}/{d["_path"].stem}.docx'
            body.append(
                f'<article id="{anchor}"><div class="meta"><span class="id">{html_mod.escape(d["id"])}</span> '
                f'<span class="type">{d["type"]}</span> <a class="dl" href="../{rel}">Download .docx</a> '
                f'<a class="dl" href="../src/{d["_path"].relative_to(SRC).as_posix()}">Markdown</a></div>'
                f'<h2>{html_mod.escape(d["title"])}</h2><p class="sum">{html_mod.escape(d.get("summary", ""))}</p>'
                f'<p class="maps">Serves: {maps}. Owner: {html_mod.escape(d["owner"])}.</p>'
                f'{md_to_html(d["_body"])}</article>')
        for r in regs:
            anchor = r["id"].lower()
            nav.append(f'<a href="#{anchor}">{html_mod.escape(r["id"])} {html_mod.escape(r["title"])} (Excel)</a>')
            cols = "".join(f"<li><b>{html_mod.escape(c['name'])}</b>{': ' + html_mod.escape(c['help']) if c.get('help') else ''}</li>" for c in r["columns"])
            body.append(
                f'<article id="{anchor}"><div class="meta"><span class="id">{html_mod.escape(r["id"])}</span> <span class="type">register</span> '
                f'<a class="dl" href="../kit/{fw}/{r["key"]}.xlsx">Download .xlsx</a></div>'
                f'<h2>{html_mod.escape(r["title"])}</h2><p class="sum">{html_mod.escape(r["purpose"])}</p>'
                f'<p class="maps">Serves: {"; ".join(html_mod.escape(m) for m in r.get("maps", []))}.</p><p>Columns:</p><ul>{cols}</ul></article>')
        nav.append("</div>")
        body.append("</section>")
    page = (DOCS / "kit_template.html").read_text(encoding="utf-8")
    page = page.replace("<!--NAV-->", "\n".join(nav)).replace("<!--BODY-->", "\n".join(body))
    page = page.replace("<!--COUNTS-->", f"{len(docs)} documents and {len(registers)} registers")
    (DOCS / "kit.html").write_text(page, encoding="utf-8")
    # landing page with live counts
    counts = {fw: sum(1 for d in docs if d["framework"] == fw) + sum(1 for r in registers if r["framework"] == fw) for fw in FRAMEWORKS}
    idx = (DOCS / "index_template.html").read_text(encoding="utf-8")
    idx = idx.replace("<!--N_DOCS-->", str(len(docs))).replace("<!--N_REGS-->", str(len(registers)))
    for fw, n in counts.items():
        idx = idx.replace(f"<!--N_{fw}-->", str(n))
    (DOCS / "index.html").write_text(idx, encoding="utf-8")


# ----------------------------------------------------------------------------- README index
def build_index(docs, registers):
    lines = []
    for fw, (name, desc) in FRAMEWORKS.items():
        group = [d for d in docs if d["framework"] == fw]
        regs = [r for r in registers if r["framework"] == fw]
        if not group and not regs:
            continue
        lines.append(f"\n### {name}\n\n{desc}.\n\n| ID | Document | Type | Serves | Download |\n|---|---|---|---|---|")
        for d in group:
            rel = f'kit/{fw}/{d["_path"].stem}.docx'
            src = f'src/{d["_path"].relative_to(SRC).as_posix()}'
            lines.append(f'| {d["id"]} | [{d["title"]}]({src}) | {d["type"]} | {"; ".join(d.get("maps", []))} | [.docx]({rel}) |')
        for r in regs:
            lines.append(f'| {r["id"]} | {r["title"]} | register | {"; ".join(r.get("maps", []))} | [.xlsx](kit/{fw}/{r["key"]}.xlsx) |')
    readme = ROOT / "README.md"
    text = readme.read_text(encoding="utf-8")
    text = re.sub(r"documents-(\d+|<!--N_DOCS-->)-blue", f"documents-{len(docs)}-blue", text)
    text = re.sub(r"registers-(\d+|<!--N_REGS-->)-blue", f"registers-{len(registers)}-blue", text)
    new = re.sub(r"(<!-- INDEX:START -->).*?(<!-- INDEX:END -->)", lambda m: m.group(1) + "\n" + "\n".join(lines) + "\n\n" + m.group(2), text, flags=re.S)
    readme.write_text(new, encoding="utf-8")


def load_registers():
    regs = []
    for p in sorted((SRC / "registers").glob("*.yaml")):
        regs.extend(yaml.safe_load(p.read_text(encoding="utf-8")))
    return regs


def main():
    docs, registers = load_docs(), load_registers()
    if "--index" not in sys.argv:
        for d in docs:
            build_docx(d)
        build_site(docs, registers)
    build_index(docs, registers)
    manifest = {d["id"]: {"title": d["title"], "type": d["type"], "framework": d["framework"], "file": d["_path"].relative_to(ROOT).as_posix()} for d in docs}
    manifest.update({r["id"]: {"title": r["title"], "type": "register", "framework": r["framework"], "file": f'kit/{r["framework"]}/{r["key"]}.xlsx'} for r in registers})
    (ROOT / "manifest.json").write_text(json.dumps(manifest, indent=1), encoding="utf-8")
    print(f"built {len(docs)} documents and indexed {len(registers)} registers on {date.today()}")


if __name__ == "__main__":
    main()

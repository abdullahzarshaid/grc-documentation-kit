"""Build Excel registers from src/registers/*.yaml into kit/<framework>/<key>.xlsx.

Layout matches the published GRC starter trackers: title block in rows 2 to 9, header in row 10,
100 prepared input rows (11 to 110) inside an Excel table, pale-yellow input cells, blue
calculated cells, dropdown validation, frozen header. One sheet per workbook, no macros.
"""
import sys
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
SRC, KIT = ROOT / "src" / "registers", ROOT / "kit"
ARIAL = "Arial"
INK, MUTED, TITLE = "FF17263A", "FF526478", "FF1661A2"
HEAD_FILL, INPUT_FILL, CALC_FILL, NOTE_FILL = "FF172D4A", "FFFFFBEF", "FFEAF2FB", "FFFFF5DA"
FIRST = 11


def as_text(v):
    if v is True: return "Yes"
    if v is False: return "No"
    return v


def build(reg):
    rows = [[as_text(v) for v in row] for row in (reg.get("rows") or [])]
    LAST = max(110, 10 + len(rows) + 10)
    wb = Workbook(); ws = wb.active; ws.title = (reg.get("sheet") or reg["title"])[:31]
    cols = reg["columns"]; ncol = len(cols); last_col = get_column_letter(ncol)
    ws.sheet_properties.tabColor = TITLE[2:]
    # title block
    ws.merge_cells(f"A2:{last_col}2"); ws["A2"] = reg["title"]
    ws["A2"].font = Font(name=ARIAL, size=16, bold=True, color=TITLE)
    prepared = f"{len(rows)} prepared rows plus {LAST - 10 - len(rows)} blank rows" if rows else "100 prepared rows"
    ws["A3"] = f"{reg['id']}. {prepared}. Keep completed copies in your organisation's approved storage; do not publish filled registers."
    ws["A3"].font = Font(name=ARIAL, size=10, color=MUTED)
    ws["A4"] = "Serves: " + "; ".join(reg.get("maps", [])); ws["A4"].font = Font(name=ARIAL, size=10, color=MUTED)
    labels = [("Recorded items", f"=COUNTA(A{FIRST}:A{LAST})")]
    status_col = next((i for i, c in enumerate(cols) if c["name"].lower().startswith("status")), None)
    if status_col is not None:
        L = get_column_letter(status_col + 1)
        labels.append(("Statuses entered", f"=COUNTA({L}{FIRST}:{L}{LAST})"))
    labels.append(("Scope / organisation", None))
    for i, (lab, formula) in enumerate(labels):
        c = ws.cell(5, i + 1, lab); c.font = Font(name=ARIAL, size=10, color=INK)
        v = ws.cell(6, i + 1, formula); v.font = Font(name=ARIAL, size=10, color=INK); v.alignment = Alignment(vertical="center")
        if formula is None:
            v.fill = PatternFill("solid", fgColor=NOTE_FILL)
    ws.merge_cells(f"A8:{last_col}8"); ws["A8"] = reg["purpose"]
    ws["A8"].font = Font(name=ARIAL, size=10, color=INK); ws["A8"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[8].height = 44
    ws.merge_cells(f"A9:{last_col}9"); ws["A9"] = "Editable inputs are pale yellow. Calculated cells are pale blue and should not be typed over. Add rows inside the table and check formulas after inserting."
    ws["A9"].font = Font(name=ARIAL, size=10, color=MUTED)
    # header
    for i, c in enumerate(cols, start=1):
        h = ws.cell(10, i, c["name"]); h.font = Font(name=ARIAL, size=10, bold=True, color="FFFFFFFF")
        h.fill = PatternFill("solid", fgColor=HEAD_FILL); h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = c.get("width", 18)
    ws.row_dimensions[10].height = 34
    # body
    for r in range(FIRST, LAST + 1):
        ws.row_dimensions[r].height = 28
        for i, c in enumerate(cols, start=1):
            cell = ws.cell(r, i); cell.font = Font(name=ARIAL, size=10, color=INK)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if c.get("formula"):
                cell.value = c["formula"].replace("{r}", str(r)); cell.fill = PatternFill("solid", fgColor=CALC_FILL)
                if c.get("number_format"): cell.number_format = c["number_format"]
            else:
                cell.fill = PatternFill("solid", fgColor=INPUT_FILL)
                if c.get("number_format"): cell.number_format = c["number_format"]
                ri = r - FIRST
                if ri < len(rows) and i - 1 < len(rows[ri]) and rows[ri][i - 1] not in (None, ""):
                    cell.value = rows[ri][i - 1]
                    if c.get("fixed"):
                        cell.fill = PatternFill("solid", fgColor="FFF4F6F8")
    # validations
    for i, c in enumerate(cols, start=1):
        L = get_column_letter(i); rng = f"{L}{FIRST}:{L}{LAST}"
        if c.get("values"):
            dv = DataValidation(type="list", formula1='"' + ",".join(str(as_text(v)) for v in c["values"]) + '"', allow_blank=True, showDropDown=False)
            dv.error, dv.errorTitle = "Choose a value from the list.", "Invalid entry"
            ws.add_data_validation(dv); dv.add(rng)
        elif c.get("range"):
            lo, hi = c["range"]
            dv = DataValidation(type="whole", operator="between", formula1=str(lo), formula2=str(hi), allow_blank=True)
            dv.error, dv.errorTitle = f"Enter a whole number from {lo} to {hi}.", "Invalid entry"
            ws.add_data_validation(dv); dv.add(rng)
        elif c.get("date"):
            dv = DataValidation(type="date", operator="greaterThan", formula1="1", allow_blank=True)
            dv.error, dv.errorTitle = "Enter a date.", "Invalid entry"
            ws.add_data_validation(dv); dv.add(rng)
    tname = "".join(ch for ch in reg["key"].title() if ch.isalnum()) + "Table"
    t = Table(displayName=tname, ref=f"A10:{last_col}{LAST}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)
    ws.freeze_panes = f"B{FIRST}"
    ws.sheet_view.showGridLines = False
    out = KIT / reg["framework"] / f"{reg['key']}.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True); wb.save(out)
    return out


def main():
    n = 0
    for p in sorted(SRC.glob("*.yaml")):
        for reg in yaml.safe_load(p.read_text(encoding="utf-8")):
            build(reg); n += 1
    print(f"built {n} registers")


if __name__ == "__main__":
    main()
